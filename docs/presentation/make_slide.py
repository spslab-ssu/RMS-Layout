# -*- coding: utf-8 -*-
"""경진대회용 '02 연구 목적' 슬라이드 (Saffar 논문 설명 + 한계점). 사용자 템플릿 배치를 재현."""
import sys
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

OUT = sys.argv[1]
FONT = "맑은 고딕"
MATH = "Cambria"

def rgb(h): return RGBColor.from_string(h)
SLATE, SLATE_D, LAV = "70798F", "667079", "D5DBEC"
BOX2, BORDER2 = "F6F8F9", "AAB8DA"
INK, BODY, MUTED = "2E3A59", "3A4458", "6B7385"
ACCENT, ACCENT_L = "C44536", "F2C9C3"
NEUTRAL, NEUTRAL_L = "8E97AD", "D5DBEC"
CARD_BORDER, FORMULA_BG, CONCL = "DCE1EC", "EEF1F8", "4A5470"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes
PX = 13.333 / 1782  # 템플릿 스크린샷 px -> inch


def rect(x, y, w, h, fill=None, line=None, lw=None, shape=MSO_SHAPE.RECTANGLE, dash=None, name=None):
    s = S.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw or 1)
        if dash: s.line.dash_style = dash
    else: s.line.fill.background()
    s.shadow.inherit = False
    if name: s.name = name
    return s


def poly(points_px, fill, name=None):
    pts = [(Emu(int(Inches(x * PX))), Emu(int(Inches(y * PX)))) for x, y in points_px]
    fb = S.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=True)
    s = fb.convert_to_shape()
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    if name: s.name = name
    return s


def style_run(r, size, bold=False, color=BODY, italic=False, font=FONT, sub=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None: el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
    rPr.set("lang", "ko-KR"); rPr.set("altLang", "en-US")
    if sub: rPr.set("baseline", "-25000")


def bulletize(p, indent=0.17):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent)))); pPr.set("indent", str(-int(Inches(indent))))
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•")


def text(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margin=0.0, name=None, space_after=0):
    """paras: list of paragraphs; each paragraph = dict(runs=[(text, opts)], bullet=bool, align=..)"""
    tb = S.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        if space_after or para.get("space_after"): p.space_after = Pt(para.get("space_after", space_after))
        p.line_spacing = para.get("line_spacing", 1.0)
        pPr = p._p.get_or_add_pPr(); pPr.set("eaLnBrk", "0"); pPr.set("latinLnBrk", "0")
        for t, o in para["runs"]:
            r = p.add_run(); r.text = t; style_run(r, **o)
        if para.get("bullet"): bulletize(p, para.get("indent", 0.17))
    if name: tb.name = name
    return tb


def P(*runs, **kw):
    return dict(runs=list(runs), **kw)


# ───────────────────────── 템플릿 프레임 ─────────────────────────
rect(0, 37 * PX, 97 * PX, 56 * PX, fill="595959", name="title-tab")
text(0, 37 * PX, 97 * PX, 56 * PX, [P(("02", dict(size=24, bold=True, color="FFFFFF", font="Arial")))],
     anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, name="title-num")
text(0.80, 0.16, 5.0, 0.66, [P(("연구 목적", dict(size=34, bold=True, color="3B3B3B")))],
     anchor=MSO_ANCHOR.MIDDLE, name="title")

def section_tab(y_px, right_px, label, name):
    rect(86 * PX, y_px * PX, (right_px - 86) * PX, 49 * PX, fill=SLATE, name=name)
    poly([(right_px - 60, y_px), (right_px, y_px), (right_px, y_px + 49), (right_px - 120, y_px + 49)], SLATE_D, name=name + "-diag")
    text(0.76, y_px * PX, 3.6, 49 * PX, [P((label, dict(size=18, bold=True, color="FFFFFF")))],
         anchor=MSO_ANCHOR.MIDDLE, name=name + "-label")

section_tab(126, 607, "1. Saffar 논문 설명", "sec1-tab")
B1X, B1Y, B1W, B1H = 86 * PX, 175 * PX, (1720 - 86) * PX, (367 - 175) * PX
rect(B1X, B1Y, B1W, B1H, fill=LAV, name="sec1-box")
section_tab(400, 457, "2. 한계점", "sec2-tab")
B2X, B2Y, B2W, B2H = 86 * PX, 449 * PX, (1720 - 86) * PX, (987 - 449) * PX
rect(B2X, B2Y, B2W, B2H, fill=BOX2, line=BORDER2, lw=3, name="sec2-box")

# ───────────────────────── 1. Saffar 논문 설명 ─────────────────────────
text(B1X + 0.22, B1Y + 0.10, 7.75, B1H - 0.2, [
    P(("Saffar, Moghaddam & Huang (2025), IJPR", dict(size=14, bold=True, color=INK)),
      ("  —  Dynamic layout design for scalable RMS", dict(size=12, italic=True, color=MUTED)), space_after=4),
    P(("재구성 가능 기계(RMT)의 ", dict(size=14)), ("구매·배치·재구성·공정 배정", dict(size=14, bold=True, color=INK)),
      ("을 하나의 MILP로 동시 결정", dict(size=14)), bullet=True, space_after=3),
    P(("목적함수: 기계 구매비 + 재구성비 + 자재취급비(흐름량 × 거리) 최소화", dict(size=14)), bullet=True, space_after=3),
    P(("수요 변화에는 1기 배치를 유지한 채 ", dict(size=14)), ("모듈 재구성으로만 대응", dict(size=14, bold=True, color=INK)),
      ("  (Robust layout)", dict(size=14, color=MUTED)), bullet=True),
], anchor=MSO_ANCHOR.MIDDLE, name="sec1-text")

# 입력 → MILP → 출력
fx, fy, fw, fh, gap = 8.82, B1Y + 0.17, 1.17, B1H - 0.34, 0.27
flow = [
    ("입력", ["기간별 수요", "공정 순서", "RMT 구성·능력"], False),
    ("MILP", ["논문 식 (1)–(16)", "총비용 최소화"], True),
    ("출력", ["구매·설치 위치", "공정 배정", "기간별 재구성"], False),
]
for i, (head, lines, dark) in enumerate(flow):
    bx = fx + i * (fw + gap)
    rect(bx, fy, fw, fh, fill=SLATE if dark else "FFFFFF", line=None if dark else SLATE, lw=1.25,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"flow-{i}").adjustments[0] = 0.12
    col = "FFFFFF" if dark else INK
    text(bx + 0.04, fy + 0.04, fw - 0.08, fh - 0.08,
         [P((head, dict(size=12.5, bold=True, color=col)), align=PP_ALIGN.CENTER, space_after=2)] +
         [P((ln, dict(size=10.5, color="FFFFFF" if dark else BODY)), align=PP_ALIGN.CENTER) for ln in lines],
         anchor=MSO_ANCHOR.MIDDLE, name=f"flow-{i}-text")
    if i < 2:
        rect(bx + fw + 0.05, fy + fh / 2 - 0.09, gap - 0.10, 0.18, fill=SLATE, shape=MSO_SHAPE.RIGHT_ARROW, name=f"flow-arrow-{i}")

# ───────────────────────── 2. 한계점: 카드 3장 ─────────────────────────
PAD, CGAP = 0.18, 0.20
CW = (B2W - 2 * PAD - 2 * CGAP) / 3
CY = B2Y + 0.17
CONC_H, CONC_GAP = 0.42, 0.12
CH = (B2Y + B2H - 0.13 - CONC_H - CONC_GAP) - CY
cards_x = [B2X + PAD + i * (CW + CGAP) for i in range(3)]
titles = ["레이아웃 위치가 전 기간 고정", "최적해 계산 부담이 큼", "제한적 예제로 일반화 어려움"]
for i, (cx, t) in enumerate(zip(cards_x, titles)):
    rect(cx, CY, CW, CH, fill="FFFFFF", line=CARD_BORDER, lw=1, name=f"card-{i+1}")
    rect(cx + 0.15, CY + 0.13, 0.34, 0.34, fill=ACCENT if i == 0 else SLATE, shape=MSO_SHAPE.OVAL, name=f"card-{i+1}-badge")
    text(cx + 0.15, CY + 0.13, 0.34, 0.34, [P((str(i + 1), dict(size=14, bold=True, color="FFFFFF", font="Arial")))],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, name=f"card-{i+1}-num")
    text(cx + 0.58, CY + 0.10, CW - 0.68, 0.40, [P((t, dict(size=15, bold=True, color=INK)))],
         anchor=MSO_ANCHOR.MIDDLE, name=f"card-{i+1}-title")
# 배지 색은 모두 같은 slate로 통일
for shp in S:
    if shp.name.endswith("-badge"):
        shp.fill.fore_color.rgb = rgb(SLATE)

IX = 0.15            # 카드 안쪽 여백
IW = CW - 2 * IX

# ── 카드 1: 위치 고정 (수식 + 미니 레이아웃) ──
x0, y0 = cards_x[0], CY
rect(x0 + IX, y0 + 0.56, IW, 0.46, fill=FORMULA_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="c1-formula-bg").adjustments[0] = 0.15
m = dict(size=17, font=MATH, italic=True, color=INK)
ms = dict(size=17, font=MATH, italic=True, color=INK, sub=True)
mr = dict(size=17, font=MATH, color=INK)
xa = dict(size=17, font=MATH, italic=True, color=ACCENT, bold=True)
xas = dict(size=17, font=MATH, italic=True, color=ACCENT, bold=True, sub=True)
text(x0 + IX, y0 + 0.56, IW, 0.46, [P(
    ("(2)   ", dict(size=12, color=MUTED)),
    ("s", m), ("pjlt", ms), ("  ≤  Σ", mr), ("j′", ms), ("Σ", mr), ("l′", ms), ("  ", mr),
    ("x", xa), ("pj′l′", xas), align=PP_ALIGN.CENTER)],
    anchor=MSO_ANCHOR.MIDDLE, name="c1-formula")
text(x0 + IX, y0 + 1.04, IW, 0.40, [
    P(("구매·설치 변수 x에 기간 인덱스 t가 없음", dict(size=11, color=ACCENT)), align=PP_ALIGN.CENTER),
    P(("→ 1기에 설치한 위치에 전 기간 고정", dict(size=11, bold=True, color=ACCENT)), align=PP_ALIGN.CENTER)],
    anchor=MSO_ANCHOR.MIDDLE, name="c1-callout")

# 미니 레이아웃 (증명된 최적해 35,775, 4x5 슬롯) — 기간 1 vs 기간 4
EMPTY = {5, 18}
RECONF = {4, 6, 7, 15, 19, 20}   # 기간 1 대비 기간 4에 구성이 바뀐 위치 (재구성 6회)
cell, cg = 0.15, 0.03
gw = 5 * cell + 4 * cg
gh = 4 * cell + 3 * cg
g1x = x0 + IX + 0.30
g2x = x0 + CW - IX - 0.30 - gw
gy = y0 + 1.68
for gi, (gx, label, show_reconf) in enumerate([(g1x, "기간 1", False), (g2x, "기간 4", True)]):
    text(gx - 0.2, gy - 0.22, gw + 0.4, 0.2, [P((label, dict(size=10.5, bold=True, color=INK)), align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE, name=f"c1-grid{gi}-label")
    for p in range(1, 21):
        col, row_b = (p - 1) % 5, (p - 1) // 5
        cx_, cy_ = gx + col * (cell + cg), gy + (3 - row_b) * (cell + cg)
        if p in EMPTY:
            rect(cx_, cy_, cell, cell, fill="FFFFFF", line=NEUTRAL, lw=0.75, dash=MSO_LINE_DASH_STYLE.DASH, name=f"c1-g{gi}-p{p}")
        else:
            rect(cx_, cy_, cell, cell, fill=ACCENT if (show_reconf and p in RECONF) else NEUTRAL, name=f"c1-g{gi}-p{p}")
mid_x = g1x + gw
text(mid_x, gy + gh / 2 - 0.22, g2x - mid_x, 0.24, [P(("→", dict(size=16, bold=True, color=INK)), align=PP_ALIGN.CENTER)],
     anchor=MSO_ANCHOR.MIDDLE, name="c1-arrow")
text(mid_x - 0.05, gy + gh / 2 + 0.0, g2x - mid_x + 0.1, 0.22, [P(("위치 동일", dict(size=10, color=MUTED)), align=PP_ALIGN.CENTER)],
     anchor=MSO_ANCHOR.MIDDLE, name="c1-same")
# 범례
ly = gy + gh + 0.07
lx = x0 + IX + 0.08
for (fill, dash, lab, w) in [(NEUTRAL, None, "고정된 기계", 1.05), (ACCENT, None, "재구성된 기계", 1.2), ("FFFFFF", True, "빈 자리", 0.9)]:
    rect(lx, ly + 0.035, 0.12, 0.12, fill=fill, line=NEUTRAL if dash else None, lw=0.75,
         dash=MSO_LINE_DASH_STYLE.DASH if dash else None, name=f"c1-legend-{lab}")
    text(lx + 0.16, ly, w - 0.16, 0.19, [P((lab, dict(size=10, color=MUTED)))], anchor=MSO_ANCHOR.MIDDLE, name=f"c1-legend-{lab}-t")
    lx += w
text(x0 + IX, y0 + 2.68, IW, CH - 2.68 - 0.06, [
    P(("예제 2 최적해: 수요가 바뀌어도(A 단종, C 신규) 18대 위치 그대로, ", dict(size=12)),
      ("재구성 6회로만 대응", dict(size=12, bold=True, color=INK)))],
    anchor=MSO_ANCHOR.TOP, name="c1-text")

# ── 카드 2: 최적 증명 시간 (gap 궤적) ──
x0 = cards_x[1]
cd = XyChartData()
ser = cd.add_series("남은 최적성 gap")
traj = [(1, 1.30), (8, 0.67), (22, 0.39), (35, 0.23), (45, 0.13), (48.1, 0.0)]
for tx, ty in traj: ser.add_data_point(tx, ty)
text(x0 + IX, y0 + 0.52, IW, 0.20, [P(("남은 최적성 gap (풀이 시간별)", dict(size=10, bold=True, color=MUTED)))],
     anchor=MSO_ANCHOR.MIDDLE, name="c2-chart-title")
gf = S.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x0 + 0.05), Inches(y0 + 0.70), Inches(CW - 0.10), Inches(1.24), cd)
gf.name = "c2-chart"
ch = gf.chart
ch.has_legend = False
ch.font.size = Pt(9); ch.font.color.rgb = rgb(MUTED); ch.font.name = FONT
xax, yax = ch.category_axis, ch.value_axis
xax.minimum_scale, xax.maximum_scale, xax.major_unit = 0, 50, 10
yax.minimum_scale, yax.maximum_scale, yax.major_unit = 0, 1.5, 0.5
xax.tick_labels.number_format, xax.tick_labels.number_format_is_linked = '0"h"', False
yax.tick_labels.number_format, yax.tick_labels.number_format_is_linked = '0.0"%"', False
yax.has_major_gridlines = True; yax.major_gridlines.format.line.color.rgb = rgb("E3E6EE")
xax.has_major_gridlines = False
for ax in (xax, yax):
    ax.format.line.color.rgb = rgb("C9CFDD"); ax.tick_labels.font.size = Pt(9); ax.tick_labels.font.color.rgb = rgb(MUTED)
s0 = ch.series[0]; s0.smooth = False
s0.format.line.color.rgb = rgb(ACCENT); s0.format.line.width = Pt(2.25)
s0.marker.style = XL_MARKER_STYLE.CIRCLE; s0.marker.size = 6
s0.marker.format.fill.solid(); s0.marker.format.fill.fore_color.rgb = rgb(ACCENT); s0.marker.format.line.color.rgb = rgb(ACCENT)
for idx, lab, pos in [(0, "1.30%", XL_LABEL_POSITION.RIGHT), (5, "0%", XL_LABEL_POSITION.ABOVE)]:
    dl = s0.points[idx].data_label
    dl.has_text_frame = True
    dl.text_frame.text = lab
    for r in dl.text_frame.paragraphs[0].runs: style_run(r, 9, bold=True, color=ACCENT)
    dl.position = pos
# 큰 숫자 + 설명
text(x0 + IX, y0 + 1.96, 1.62, 0.42, [P(("48.1시간", dict(size=24, bold=True, color=ACCENT)))],
     anchor=MSO_ANCHOR.MIDDLE, name="c2-stat")
text(x0 + IX + 1.66, y0 + 1.96, IW - 1.66, 0.42, [
    P(("예제 2 최적 증명 소요", dict(size=10, color=MUTED))),
    P(("(탐색 노드 467만 개)", dict(size=10, color=MUTED)))],
    anchor=MSO_ANCHOR.MIDDLE, name="c2-stat-caption")
text(x0 + IX, y0 + 2.43, IW, CH - 2.43 - 0.06, [
    P(("10분 풀이 시 gap 2.1% — 최적 보장 불가", dict(size=12)), bullet=True, space_after=2),
    P(("LP 완화 gap 8.4% (32,776 vs 35,775)", dict(size=12)), bullet=True, space_after=2),
    P(("논문 보고값 36,006 > ", dict(size=12)), ("증명된 최적 35,775", dict(size=12, bold=True, color=INK)), bullet=True),
], name="c2-text")

# ── 카드 3: 예제 구조의 한계 (범위 막대) ──
x0 = cards_x[2]
text(x0 + IX, y0 + 0.55, IW, 0.40, [
    P(("기계 이동 허용 시 비용 절감률", dict(size=10.5, bold=True, color=INK)), ("  (이동비 0 기준)", dict(size=10, color=MUTED))),
    P(("같은 조건(3×4 배치·수요 0.5배)에서 공정 순서만 다름", dict(size=10, color=MUTED)))],
    anchor=MSO_ANCHOR.MIDDLE, name="c3-subtitle")
scale = 2.55 / 4.5   # 1%p 당 inch
bars = [("공정 공유형 (Saffar 예제의 공정 순서)", 0.97, 1.48, NEUTRAL, NEUTRAL_L, "0.97~1.48%", INK),
        ("공정 비공유형 (부품별 전용 공정)", 2.99, 4.04, ACCENT, ACCENT_L, "2.99~4.04%", ACCENT)]
by = y0 + 1.02
for k, (lab, lo, hi, c_solid, c_light, vlab, vcol) in enumerate(bars):
    text(x0 + IX, by, IW, 0.20, [P((lab, dict(size=10.5, color=BODY)))], anchor=MSO_ANCHOR.MIDDLE, name=f"c3-bar{k}-label")
    bx0 = x0 + IX
    rect(bx0, by + 0.23, lo * scale, 0.25, fill=c_solid, name=f"c3-bar{k}-proven")
    rect(bx0 + lo * scale, by + 0.23, (hi - lo) * scale, 0.25, fill=c_light, name=f"c3-bar{k}-range")
    text(bx0 + hi * scale + 0.07, by + 0.20, 1.1, 0.30, [P((vlab, dict(size=12, bold=True, color=vcol)))],
         anchor=MSO_ANCHOR.MIDDLE, name=f"c3-bar{k}-value")
    by += 0.56
text(x0 + IX, by - 0.02, IW, 0.20, [P(("진한 부분: 증명된 최소 · 옅은 부분: 가능 범위", dict(size=10, color=MUTED)))],
     anchor=MSO_ANCHOR.MIDDLE, name="c3-legend")
text(x0 + IX, y0 + 2.36, IW, CH - 2.36 - 0.06, [
    P(("예제가 소규모(16·20슬롯, 4기간)이고 3부품 모두 공정 2→12 공유", dict(size=12)), bullet=True, space_after=2),
    P(("위치 고정이 전제 → ", dict(size=12)), ("재배치 가치 평가 불가", dict(size=12, bold=True, color=INK)), bullet=True),
], name="c3-text")

# ── 결론 줄 ──
cy = B2Y + B2H - 0.13 - CONC_H
rect(B2X + PAD, cy, B2W - 2 * PAD, CONC_H, fill=CONCL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="conclusion-bg").adjustments[0] = 0.25
text(B2X + PAD + 0.2, cy, B2W - 2 * PAD - 0.4, CONC_H, [P(
    ("→  본 연구 목적:  ", dict(size=14, bold=True, color="FFFFFF")),
    ("기간마다 기계 재배치를 허용하는 Adaptive layout MILP를 제안하고, 재배치가 이득이 되는 조건을 규명", dict(size=14, color="FFFFFF")))],
    anchor=MSO_ANCHOR.MIDDLE, name="conclusion")

slide.notes_slide.notes_text_frame.text = (
    "[출처·근거]\n"
    "논문: Saffar, Moghaddam & Huang (2025), Dynamic layout design for scalable reconfigurable manufacturing systems, IJPR 63(17), 6530–6552.\n"
    "한계 1: 논문 식 (2) s_pjlt ≤ Σ_j′Σ_l′ x_pj′l′. 구매 변수 x에 기간 인덱스가 없어 1기에 설치한 위치에 전 기간 결박. 식 (5)의 상태 전이도 같은 위치 p 안의 유지·재구성만 허용. "
    "미니 레이아웃은 논문 예제 2(multi_part)의 증명된 최적해 35,775: 18대, 빈 자리 5·18, 재구성 6회(위치 4·6·7·15·19·20).\n"
    "한계 2: 같은 정식화로 Gurobi, 시간 제한 없음, MIPGap 0 → 48.10시간(173,154초) 만에 최적 증명, 노드 4,667,187개, 최적해 발견 31.0시간. "
    "LP 완화 하한 32,776.4(root gap 8.4%). gap 궤적 1h 1.30% / 8h 0.67% / 22h 0.39% / 35h 0.23% / 45h 0.13% / 48.1h 0%. 10분(600초) 실행 gap 2.09%. 논문 보고값 36,006.\n"
    "한계 3: 3×4 레이아웃, 수요 0.5배, 재구성비 50/25, 이동비 0(이득 상한). 공정 공유형(원본 순서 A 2>12>17, B 2>12>11, C 2>12>11>8) 이득 구간 [173, 264] / off 17,829 → 0.97~1.48%. "
    "공정 비공유형(A 18>9>19, B 10>1>17, C 13>16>8>12) 이득 구간 [530, 717] / off 17,730 → 2.99~4.04%. 두 경우 모두 off 최적 증명.\n"
)
prs.save(OUT)
print("saved", OUT)
